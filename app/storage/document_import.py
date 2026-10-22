"""使用專案 Python 解析套件擷取文字；不依賴 Office、系統轉檔程式或模型。"""
import io
import json
import re
import zipfile
from pathlib import Path
from uuid import uuid4
from bs4 import BeautifulSoup
from app.logging_setup import log
from app.storage.workspace import resolve_path

MAX_UPLOAD = 25 * 1024 * 1024
IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg", ".ico", ".heic", ".avif"}
BINARY = {".exe", ".dll", ".zip", ".7z", ".rar", ".mp3", ".wav", ".mp4", ".mov", ".avi", ".db", ".sqlite", ".pyc", ".bin"}
SUPPORTED = "PDF、DOCX、PPTX、XLSX、XLS、ODT/ODS/ODP、EPUB、RTF、HTML、TXT/Markdown/CSV/JSON/XML 與程式碼等純文字"
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".xls", ".odt", ".ods", ".odp", ".epub", ".rtf", ".doc", ".ppt", ".docm", ".pptm"}


def decode_text(data):
    if data.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return data.decode("utf-32")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if b"\0" in data:
        raise ValueError("此檔案不是可直接讀取的文字。")
    for encoding in ("utf-8-sig", "cp950", "cp1252"):
        try:
            text = data.decode(encoding)
            controls = sum(ord(char) < 32 and char not in "\n\r\t\f" for char in text)
            if controls > max(1, len(text) // 100):
                break
            return text
        except UnicodeError:
            continue
    raise ValueError("無法辨識文字編碼，請另存為 UTF-8 後匯入。")


def html_text(data):
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "img", "svg", "picture"]):
        tag.decompose()
    return soup.get_text("\n", strip=True)


def package_text(data, suffix):
    """各格式只存取套件的文字／表格 API，不轉圖片、不執行文件內的程式。"""
    stream = io.BytesIO(data)
    blocks = []
    if suffix == ".docx":
        from docx import Document
        document = Document(stream)
        def paragraphs(container):
            for item in container.iter_inner_content():
                if hasattr(item, "rows"):
                    for row in item.rows:
                        yield "\t".join(" ".join(paragraphs(cell)) for cell in row.cells)
                else:
                    yield item.text
        blocks.append(("文件內文", "\n".join(paragraphs(document))))
        for index, section in enumerate(document.sections, 1):
            for label, part in (("頁首", section.header), ("頁尾", section.footer)):
                text = "\n".join(paragraphs(part))
                if text.strip():
                    blocks.append((f"第 {index} 節 {label}", text))
    elif suffix == ".pptx":
        from pptx import Presentation
        def shape_text(shapes):
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    yield from shape_text(shape.shapes)
                if shape.has_text_frame:
                    yield shape.text_frame.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield "\t".join(cell.text for cell in row.cells)
        for index, slide in enumerate(Presentation(stream).slides, 1):
            text = "\n".join(shape_text(slide.shapes))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                text += "\n備註：" + slide.notes_slide.notes_text_frame.text
            blocks.append((f"投影片 {index}", text))
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                rows = ["\t".join("" if value is None else str(value) for value in row)
                        for row in sheet.iter_rows(values_only=True)]
                blocks.append((f"工作表 {sheet.title}", "\n".join(rows)))
        finally:
            book.close()
    elif suffix in {".odt", ".ods", ".odp"}:
        from odf.opendocument import load
        from odf.teletype import extractText
        from odf.namespaces import TEXTNS
        def paragraphs(node):
            if getattr(node, "qname", None) in {(TEXTNS, "p"), (TEXTNS, "h")}:
                yield extractText(node)
            else:
                for child in getattr(node, "childNodes", []):
                    yield from paragraphs(child)
        blocks.append(("文件文字", "\n".join(paragraphs(load(stream).body))))
    elif suffix == ".epub":
        from ebooklib import epub
        book = epub.read_epub(stream, options={"ignore_ncx": True})
        for index, (identifier, _) in enumerate(book.spine, 1):
            item = book.get_item_with_id(identifier)
            if item:
                blocks.append((f"章節 {index} · {item.get_name()}", html_text(item.get_content())))
    return blocks


def extract_text(data, name):
    suffix = Path(name).suffix.lower()
    image_magic = data.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF87a", b"GIF89a", b"BM", b"II*\0", b"MM\0*")) or (data.startswith(b"RIFF") and data[8:12] == b"WEBP")
    if suffix in IMAGES or image_magic:
        raise ValueError("不匯入圖片；此功能只擷取文件中已存在的文字，不做 OCR。")
    warnings = []
    if suffix == ".pdf" or data.startswith(b"%PDF-"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF 有密碼保護，請先解除保護後匯入。")
        if len(reader.pages) > 1000:
            raise ValueError("PDF 超過 1000 頁，請拆分後匯入。")
        blocks = [(f"第 {index} 頁", page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
        empty = sum(not text.strip() for _, text in blocks)
        if empty:
            warnings.append(f"{empty} 頁沒有可擷取文字；圖片與掃描影像不做辨識。")
    elif suffix in {".docx", ".pptx", ".xlsx", ".xlsm", ".odt", ".ods", ".odp", ".epub"}:
        if zipfile.is_zipfile(io.BytesIO(data)):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist() if item.filename.endswith((".xml", ".html", ".xhtml"))) > 100 * 1024 * 1024:
                    raise ValueError("文件內文字資料過大，請拆分後匯入。")
        blocks = package_text(data, suffix)
    elif suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            blocks = [(f"工作表 {sheet.name}", "\n".join("\t".join(str(value) for value in sheet.row_values(row))
                       for row in range(sheet.nrows))) for sheet in book.sheets()]
        finally:
            book.release_resources()
    elif suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text
        blocks = [("文件文字", rtf_to_text(data.decode("latin-1")))]
    elif suffix in {".doc", ".ppt", ".docm", ".pptm"}:
        raise ValueError("此 Word/PowerPoint 格式請先另存為 DOCX/PPTX；專案不依賴系統 Office 轉檔程式。")
    elif suffix in BINARY or data.startswith((b"MZ", b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
        raise ValueError("此二進位格式尚無文字解析器；請另存為支援的文件或文字格式。")
    else:
        text = html_text(data) if suffix in {".html", ".htm", ".xhtml"} else decode_text(data)
        blocks = [("文件文字", text)]
    meaningful = any(text.strip() for _, text in blocks)
    if not meaningful:
        warnings.append("文件沒有可擷取的文字；未呼叫模型或進行圖片辨識。")
    text = "\n\n".join(f"## {label}\n\n{content}" for label, content in blocks if content.strip())
    return text, warnings


def import_document(data, name):
    name = Path(name.replace("\\", "/")).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .") or "document.txt"
    suffix = Path(name).suffix[:16]
    name = Path(name).stem[:80 - len(suffix)] + suffix
    if name.split(".", 1)[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        name = "file_" + name
    if len(data) > MAX_UPLOAD:
        raise ValueError("單一文件最多 25 MiB。")
    text, warnings = extract_text(data, name)
    relative = f"documents/{uuid4().hex}"
    folder = Path(resolve_path(relative))
    folder.mkdir(parents=True, exist_ok=False)
    (folder / "source").mkdir()
    (folder / "source" / name).write_bytes(data)
    text_name = "extracted-text.md"
    text_path = f"{relative}/{text_name}" if text else None
    if text:
        (folder / text_name).write_text(f"# 來源：{name}\n\n{text}\n", encoding="utf-8")
    result = {"name": name, "original_path": f"{relative}/source/{name}", "text_path": text_path,
              "characters": len(text), "warnings": warnings, "status": "imported" if text else "no_text"}
    (folder / "import.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    log("document.imported", **result)
    return result


def read_workspace_text(path):
    """程式與純文字保持原始行號；結構化文件才透過解析套件轉成文字。"""
    path = Path(path)
    if path.stat().st_size > MAX_UPLOAD:
        raise ValueError("文件超過 25 MiB，請拆分後閱讀。")
    if path.suffix.lower() in IMAGES:
        raise ValueError("不讀取圖片；僅支援文件文字。")
    data = path.read_bytes()
    if path.suffix.lower() in DOCUMENT_SUFFIXES or data.startswith(b"%PDF-"):
        text, warnings = extract_text(data, path.name)
        return text + ("\n\n解析說明：" + "；".join(warnings) if warnings else "")
    return decode_text(data)
